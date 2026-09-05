from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUTPUT = Path(r"C:\laragon\www\cbta_app\output\Informe_desarrollos_05ago_01sep_2026.docx")

BLUE = "2E74B5"
DARK_BLUE = "1F4D78"
LIGHT_BLUE = "E8EEF5"
LIGHT_GRAY = "F2F4F7"
INK = "0B2545"


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_width(cell, width_dxa):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(width_dxa))
    tc_w.set(qn("w:type"), "dxa")


def keep_row_together(row):
    tr = getattr(row, "_tr", row)
    tr_pr = tr.get_or_add_trPr()
    if tr_pr.find(qn("w:cantSplit")) is None:
        tr_pr.append(OxmlElement("w:cantSplit"))


def set_table_geometry(table, widths):
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table_pr = table._tbl.tblPr
    tbl_w = table_pr.find(qn("w:tblW"))
    tbl_w.set(qn("w:w"), "9360")
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = OxmlElement("w:tblInd")
    tbl_ind.set(qn("w:w"), "120")
    tbl_ind.set(qn("w:type"), "dxa")
    table_pr.append(tbl_ind)
    for row in table.rows:
        keep_row_together(row)
        for cell, width in zip(row.cells, widths):
            set_cell_width(cell, width)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_after = Pt(2)
                paragraph.paragraph_format.space_before = Pt(2)


def set_font(run, size=11, bold=False, color=INK):
    run.font.name = "Calibri"
    run._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    run._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)


def paragraph(doc, text="", style=None, bold_prefix=None):
    p = doc.add_paragraph(style=style)
    if bold_prefix and text.startswith(bold_prefix):
        r = p.add_run(bold_prefix)
        set_font(r, bold=True)
        r = p.add_run(text[len(bold_prefix):])
        set_font(r)
    else:
        r = p.add_run(text)
        set_font(r)
    return p


def bullet(doc, text):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run(text)
    set_font(r)
    return p


def heading(doc, text, level=1):
    p = doc.add_paragraph(style=f"Heading {level}")
    r = p.add_run(text)
    set_font(r, size={1: 16, 2: 13, 3: 12}[level], bold=True, color=BLUE if level < 3 else DARK_BLUE)
    return p


def add_meta_table(doc):
    table = doc.add_table(rows=4, cols=2)
    set_table_geometry(table, [2700, 6660])
    rows = [
        ("Periodo evaluado", "5 de agosto al 1 de septiembre de 2026"),
        ("Tarifa acordada", "$100.00 MXN por hora"),
        ("Estimación recomendada", "80 horas - $8,000.00 MXN"),
        ("Alcance excluido", "API/Dr. Sam, aplicación móvil y funcionalidad original de ramas ggh"),
    ]
    for row, (label, value) in zip(table.rows, rows):
        set_cell_shading(row.cells[0], LIGHT_BLUE)
        label_run = row.cells[0].paragraphs[0].add_run(label)
        set_font(label_run, bold=True, color=DARK_BLUE)
        value_run = row.cells[1].paragraphs[0].add_run(value)
        set_font(value_run)
    doc.add_paragraph()


def add_estimate_table(doc):
    table = doc.add_table(rows=1, cols=4)
    set_table_geometry(table, [2700, 3500, 1100, 2060])
    headers = ["Bloque", "Alcance incluido", "Horas", "Importe"]
    for cell, text in zip(table.rows[0].cells, headers):
        set_cell_shading(cell, LIGHT_BLUE)
        run = cell.paragraphs[0].add_run(text)
        set_font(run, bold=True, color=DARK_BLUE)
        cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER

    rows = [
        ("Inventario nutricional", "Lotes, existencias y movimientos.", "8", "$800"),
        ("Documentos y trazabilidad", "PDFs, QR y responsables operativos.", "8", "$800"),
        ("Cargos adicionales", "Listas de precio, aplicación automática y remisiones.", "18", "$1,800"),
        ("Inventario oncológico avanzado", "mL, remanentes, respaldo, mermas y pérdidas.", "22", "$2,200"),
        ("Catálogo y permisos", "Edición unificada y control de Super Administrador.", "7", "$700"),
        ("Integración ggh3", "Ajustes propios de rutas, imports y menús tras merge.", "3", "$300"),
        ("Preparación e inspección", "Tiempos, responsables sanitarios y flujo de calidad.", "10", "$1,000"),
        ("Pruebas y validación", "Migraciones, pruebas de flujo y verificación de vistas.", "4", "$400"),
    ]
    for values in rows:
        cells = table.add_row().cells
        keep_row_together(cells[0]._tc.getparent())
        for index, value in enumerate(values):
            run = cells[index].paragraphs[0].add_run(value)
            set_font(run)
            if index >= 2:
                cells[index].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER

    total = table.add_row().cells
    keep_row_together(total[0]._tc.getparent())
    for cell in total:
        set_cell_shading(cell, LIGHT_GRAY)
    total[0].merge(total[1])
    run = total[0].paragraphs[0].add_run("TOTAL ESTIMADO")
    set_font(run, bold=True, color=DARK_BLUE)
    total[2].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    total[3].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = total[2].paragraphs[0].add_run("80")
    set_font(run, bold=True, color=DARK_BLUE)
    run = total[3].paragraphs[0].add_run("$8,000 MXN")
    set_font(run, bold=True, color=DARK_BLUE)


def add_detail_section(doc, title, items):
    heading(doc, title, 2)
    for item in items:
        bullet(doc, item)


def build_document():
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.1

    for level, size, before, after, color in [(1, 16, 16, 8, BLUE), (2, 13, 12, 6, BLUE), (3, 12, 8, 4, DARK_BLUE)]:
        style = doc.styles[f"Heading {level}"]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor.from_string(color)
        style.font.bold = True
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)

    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = header.add_run("CBTA - Informe de desarrollos")
    set_font(run, size=9, color="5A6B7D")
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = footer.add_run("Periodo: 05 ago - 01 sep 2026")
    set_font(run, size=9, color="5A6B7D")

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.LEFT
    title.paragraph_format.space_before = Pt(0)
    title.paragraph_format.space_after = Pt(4)
    run = title.add_run("Informe de desarrollos y estimación de servicios")
    set_font(run, size=22, bold=True, color=DARK_BLUE)
    subtitle = doc.add_paragraph()
    subtitle.paragraph_format.space_after = Pt(14)
    run = subtitle.add_run("Sistema CBTA | Corte al 1 de septiembre de 2026")
    set_font(run, size=11, color="5A6B7D")

    add_meta_table(doc)
    heading(doc, "1. Objetivo y criterio de cobro", 1)
    paragraph(doc, "Este documento delimita los desarrollos realizados durante el periodo indicado y presenta una estimación comercial a una tarifa de $100.00 MXN por hora. Su propósito es evitar el cobro duplicado de funcionalidades creadas en meses anteriores o desarrolladas por terceros.")
    paragraph(doc, "La estimación se basa en el alcance funcional, cambios de base de datos, vistas, controladores, pruebas y validaciones observadas en el repositorio. No sustituye un registro horario minuto a minuto.")

    heading(doc, "2. Exclusiones expresas", 1)
    for item in [
        "Integración API/Dr. Sam, webhooks, solicitudes externas y tareas relacionadas. Ese alcance se factura por separado.",
        "Proyecto de aplicación móvil y su configuración técnica. Se excluye por haber sido acordado para pago separado.",
        "Funcionalidad desarrollada originalmente en las ramas ggh, ggh2, ggh3 y ggh5.",
        "Correcciones puntuales sin desarrollo nuevo, por ejemplo redirecciones, errores de visualización o ajustes menores de interfaz.",
        "Cambios previos al 5 de agosto de 2026.",
    ]:
        bullet(doc, item)

    heading(doc, "3. Desarrollo nuevo incluido", 1)
    add_detail_section(doc, "3.1 Inventario nutricional y trazabilidad documental", [
        "Ampliación de existencias nutricionales por lote y consulta de movimientos.",
        "Mejoras de trazabilidad operativa en orden de preparación, etiqueta y remisión oncológica.",
        "Identificación de responsables de preparación, revisión, aprobación y liberación.",
        "Etiquetas y pantallas de consulta mediante códigos QR para nutrición y oncología.",
    ])
    add_detail_section(doc, "3.2 Cargos adicionales por lista de precio", [
        "Modelo de cargos adicionales ilimitados por lista de precio y categoría.",
        "Configuración desde creación, edición y consulta de listas.",
        "Aplicación automática por solicitud en nutrición y por mezcla en oncología/antibióticos.",
        "Desglose de cargos en remisiones y sustitución del concepto heredado de servicio de mezclado.",
    ])
    add_detail_section(doc, "3.3 Inventario oncológico avanzado", [
        "Registro de movimientos de inventario en mililitros.",
        "Control de remanentes, estabilidad y caducidad de medicamentos abiertos.",
        "Uso de almacén de respaldo cuando el almacén principal no cuenta con stock.",
        "Separación entre merma de remanente y pérdida de stock; captura de frascos/mL y motivo de pérdida.",
        "Vistas de movimientos por lote y simulaciones de consumo para validación.",
    ])
    add_detail_section(doc, "3.4 Catálogo oncológico y permisos", [
        "Unificación de la edición de medicamentos genéricos y presentaciones en una sola pantalla.",
        "Control de permisos: solo Super Administrador puede editar medicamentos genéricos; el administrador conserva creación.",
    ])
    add_detail_section(doc, "3.5 Calidad, preparación e inspección", [
        "Registro permanente de fecha y hora exacta de preparación de mezcla.",
        "Cálculo y despliegue de dosis total, volumen final y concentración en la orden de preparación.",
        "Separación de responsabilidades entre quien inspecciona y quien aprueba.",
        "Selector de aprobador limitado a Responsable sanitario o Auxiliar de responsable sanitario.",
        "Ampliación del catálogo de puestos con Auxiliar de responsable sanitario.",
    ])
    add_detail_section(doc, "3.6 Trabajo propio de integración de merge", [
        "Corrección de imports/controladores para que las rutas de Distribución funcionaran después de integrar ggh3.",
        "Integración de entradas de menú para Distribución y Super Administrador.",
        "Estos puntos se incluyen por ser ajustes propios de integración, no por la funcionalidad original de la rama.",
    ])

    doc.add_page_break()
    heading(doc, "4. Merges registrados y tratamiento comercial", 1)
    table = doc.add_table(rows=1, cols=3)
    set_table_geometry(table, [1500, 3300, 4560])
    for cell, text in zip(table.rows[0].cells, ["Fecha", "Actividad", "Tratamiento en esta estimación"]):
        set_cell_shading(cell, LIGHT_BLUE)
        run = cell.paragraphs[0].add_run(text)
        set_font(run, bold=True, color=DARK_BLUE)
    merge_rows = [
        ("17 ago", "Merge de ggh a main", "No se cobra funcionalidad original de la rama."),
        ("21 ago", "Merge de ggh2 y homologaciones", "No se cobra funcionalidad original de la rama."),
        ("26-28 ago", "Merge de ggh3 a main", "Solo se incluyen 3 h de ajustes propios de integración."),
        ("31 ago", "Merge rápido de ggh5 a main", "Excluido: no hubo conflicto ni corrección de merge."),
    ]
    for values in merge_rows:
        cells = table.add_row().cells
        keep_row_together(cells[0]._tc.getparent())
        for cell, value in zip(cells, values):
            run = cell.paragraphs[0].add_run(value)
            set_font(run)

    heading(doc, "5. Estimación de horas e importe", 1)
    paragraph(doc, "La siguiente distribución resume el esfuerzo estimado únicamente del alcance incluido. Los importes están expresados en pesos mexicanos y no incluyen IVA, en caso de que aplique.")
    add_estimate_table(doc)

    heading(doc, "6. Evidencia de no duplicidad", 1)
    for item in [
        "El periodo inicia el 5 de agosto de 2026; no se incorpora trabajo previo.",
        "Se excluyeron los módulos originados en ramas ggh y la integración API/Dr. Sam.",
        "Los cambios incluidos corresponden a inventario, cargos adicionales, catálogo, permisos, documentos y trazabilidad de calidad implementados o integrados de forma directa durante el periodo.",
        "Los ajustes correctivos simples no fueron asignados a un bloque de desarrollo nuevo.",
        "El total se presenta como estimación de alcance para revisión y aprobación antes de facturar.",
    ]:
        bullet(doc, item)

    heading(doc, "7. Recomendación de cobro", 1)
    paragraph(doc, "Con una tarifa de $100.00 MXN por hora, la recomendación es facturar 80 horas, por un total de $8,000.00 MXN más IVA, si corresponde.", bold_prefix="Con una tarifa de $100.00 MXN por hora, ")
    paragraph(doc, "Si se desea una postura conservadora para negociación, puede presentarse el mismo alcance como rango de 72 a 80 horas ($7,200 a $8,000 MXN).")

    doc.core_properties.title = "Informe de desarrollos y estimación de servicios"
    doc.core_properties.subject = "CBTA - desarrollos del 5 de agosto al 1 de septiembre de 2026"
    doc.core_properties.author = "CBTA"
    doc.save(OUTPUT)


if __name__ == "__main__":
    build_document()
